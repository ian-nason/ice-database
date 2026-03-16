#!/usr/bin/env python3
"""
ICE Database Builder

Reads ICE enforcement data from two FOIA releases (2023 + 2025) of the
Deportation Data Project, normalizes schemas, deduplicates overlapping
records, and produces a queryable DuckDB database.

Data sources:
  release_2023/ - 2023 FOIA litigation release (FY2012-FY2023)
  release_2025/ - Dec 2025 settlement release (Sep 2023 - Oct 2025)

Tables produced:
  arrests        - Administrative arrests (both releases)
  detainers      - Detainer requests (2025 only)
  detentions     - Detention stays (both releases)
  removals       - Deportation/removal records (2023 only)
  rca_decisions  - Release/custody assessment decision history (2023 only)

Usage:
    uv run python build_database.py
    uv run python build_database.py --data-dir data/raw/ --output my_ice.duckdb
"""

import argparse
import re
import sys
import time
from datetime import datetime
from pathlib import Path

import duckdb
import openpyxl
import pandas as pd
from tqdm import tqdm

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

DEFAULT_OUTPUT = "ice.duckdb"

# ---------------------------------------------------------------------------
# File classification — map filenames to table names
# ---------------------------------------------------------------------------

TABLE_PATTERNS = [
    ("arrests", re.compile(r"arrest", re.I)),
    ("detainers", re.compile(r"detainer", re.I)),
    ("detentions", re.compile(r"detention", re.I)),
    ("removals", re.compile(r"removal", re.I)),
    ("rca_decisions", re.compile(r"rca", re.I)),
]

TABLE_DESCRIPTIONS = {
    "arrests": "ICE administrative arrests",
    "detainers": "Detainer requests issued to jails/prisons",
    "detentions": "Detention stays (book-in to book-out)",
    "removals": "Deportation/removal records",
    "rca_decisions": "Release/custody assessment decision history",
}

# ---------------------------------------------------------------------------
# Column harmonization — applied after snake_case normalization
# ---------------------------------------------------------------------------

COLUMN_RENAMES = {
    # Unique identifier (varies across releases and tables)
    "unique_identifier": "unique_id",
    "anonymized_identifier": "unique_id",
    "anonymized_identifer": "unique_id",  # typo in 2023 removals
    # Detention date columns (2025 appends _time)
    "stay_book_in_date_time": "stay_book_in_date",
    "book_in_date_time": "detention_book_in_date",
    "detention_book_out_date_time": "detention_book_out_date",
    # Cross-release field names
    "marital_status": "marital",
    "most_serious_conviction_msc_charge_code": "msc_charge_code",
    "eid_case_id": "case_id",
    "eid_subject_id": "subject_id",
    # Arrest creator field variants across FY sheets
    "arrest_create_by": "arrest_created_by",
    "arrested_created_by": "arrest_created_by",
}

# Dedup keys per table — for tables present in both releases
DEDUP_KEYS = {
    "arrests": ["unique_id", "apprehension_date"],
    "detentions": ["unique_id", "stay_book_in_date", "detention_facility_code"],
}

# Date column detection
DATE_RE = re.compile(r"date", re.I)


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


def classify_file(filename: str) -> str | None:
    """Return table name for a file based on filename pattern."""
    for table_name, pattern in TABLE_PATTERNS:
        if pattern.search(filename):
            return table_name
    return None


def find_header_row(filepath: Path, sheet, max_scan: int = 15) -> int:
    """Find the real header row (row with most non-null values in first N rows)."""
    df = pd.read_excel(
        str(filepath), sheet_name=sheet, header=None,
        nrows=max_scan, engine="openpyxl",
    )
    best_row = 0
    best_count = 0
    for i, row in df.iterrows():
        count = row.notna().sum()
        if count > best_count:
            best_count = count
            best_row = i
    return best_row


def normalize_columns(df: pd.DataFrame) -> pd.DataFrame:
    """Normalize column names to snake_case."""
    df.columns = [
        re.sub(r"_+", "_", re.sub(r"[^a-z0-9]", "_", c.strip().lower())).strip("_")
        for c in df.columns
    ]
    return df


def apply_renames(df: pd.DataFrame) -> pd.DataFrame:
    """Apply standard column renames for cross-release harmonization."""
    rmap = {}
    for old, new in COLUMN_RENAMES.items():
        if old in df.columns and new not in df.columns:
            rmap[old] = new
    if rmap:
        df = df.rename(columns=rmap)
    return df


def cast_dates(df: pd.DataFrame) -> pd.DataFrame:
    """Convert date-like columns to datetime."""
    for col in df.columns:
        if DATE_RE.search(col):
            df[col] = pd.to_datetime(df[col], errors="coerce")
    return df


def get_release_name(filepath: Path) -> str:
    """Extract release name from path (e.g. 'release_2023')."""
    for part in filepath.parts:
        if part.startswith("release_"):
            return part
    return "unknown"


# ---------------------------------------------------------------------------
# Discovery
# ---------------------------------------------------------------------------


def discover_sources(data_dir: Path) -> dict[str, list[tuple[Path, str]]]:
    """
    Scan release_* directories and classify files into tables.
    Returns {table_name: [(filepath, sheet_name), ...]}.
    """
    sources: dict[str, list[tuple[Path, str]]] = {
        t: [] for t, _ in TABLE_PATTERNS
    }

    release_dirs = sorted(
        d for d in data_dir.iterdir()
        if d.is_dir() and d.name.startswith("release_")
    )
    if not release_dirs:
        print(f"  Error: No release_* directories in {data_dir}")
        sys.exit(1)

    for rdir in release_dirs:
        xlsx_files = sorted(
            f for f in rdir.iterdir()
            if f.suffix.lower() in (".xlsx", ".xls")
        )
        print(f"  {rdir.name}: {len(xlsx_files)} files")

        for fpath in xlsx_files:
            table = classify_file(fpath.stem)
            if table is None:
                print(f"    ? {fpath.name} (unclassified)")
                continue

            wb = openpyxl.load_workbook(str(fpath), read_only=True, data_only=True)
            for sheet in wb.sheetnames:
                sources[table].append((fpath, sheet))
            wb.close()

    return sources


# ---------------------------------------------------------------------------
# Table loading
# ---------------------------------------------------------------------------


def _pandas_to_duckdb_type(series: pd.Series) -> str:
    """Map a pandas Series dtype to a DuckDB column type."""
    if pd.api.types.is_datetime64_any_dtype(series):
        return "TIMESTAMP"
    if pd.api.types.is_integer_dtype(series):
        return "BIGINT"
    if pd.api.types.is_float_dtype(series):
        return "DOUBLE"
    if pd.api.types.is_bool_dtype(series):
        return "BOOLEAN"
    return "VARCHAR"


def load_table(
    con: duckdb.DuckDBPyConnection,
    table_name: str,
    chunks: list[tuple[Path, str]],
) -> int:
    """Load all source chunks into a single DuckDB table. Returns row count."""
    if not chunks:
        return 0

    total = 0
    created = False

    for fpath, sheet in tqdm(chunks, desc=f"  {table_name}", unit="sheet"):
        try:
            hrow = find_header_row(fpath, sheet)
            df = pd.read_excel(
                str(fpath), sheet_name=sheet, header=hrow, engine="openpyxl",
            )
            if len(df) == 0:
                continue

            df = normalize_columns(df)
            df = apply_renames(df)
            df = cast_dates(df)
            df["data_source"] = get_release_name(fpath)
            df = df.dropna(axis=1, how="all")

            n = len(df)

            if not created:
                con.execute(f"DROP TABLE IF EXISTS {table_name}")
                con.execute(f"CREATE TABLE {table_name} AS SELECT * FROM df")
                created = True
            else:
                # Add columns that don't exist yet, using proper DuckDB types
                existing = {
                    r[0] for r in con.execute(f"DESCRIBE {table_name}").fetchall()
                }
                for col in df.columns:
                    if col not in existing:
                        dtype = _pandas_to_duckdb_type(df[col])
                        con.execute(
                            f'ALTER TABLE {table_name} ADD COLUMN "{col}" {dtype}'
                        )
                con.execute(
                    f"INSERT INTO {table_name} BY NAME SELECT * FROM df"
                )

            total += n
            del df

        except Exception as e:
            print(f"\n    ERROR: {fpath.name} [{sheet}]: {e}")

    return total


# ---------------------------------------------------------------------------
# Deduplication
# ---------------------------------------------------------------------------


def dedup_table(
    con: duckdb.DuckDBPyConnection, table_name: str, key_cols: list[str],
) -> int:
    """Remove cross-release duplicates, preferring release_2025. Returns rows removed."""
    existing = {r[0] for r in con.execute(f"DESCRIBE {table_name}").fetchall()}
    keys = [k for k in key_cols if k in existing]

    if not keys or "data_source" not in existing:
        return 0

    before = con.execute(f"SELECT COUNT(*) FROM {table_name}").fetchone()[0]

    key_list = ", ".join(f'"{k}"' for k in keys)
    null_check = f'"{keys[0]}" IS NOT NULL'

    con.execute(f"""
        CREATE OR REPLACE TABLE {table_name} AS
        WITH ranked AS (
            SELECT *,
                ROW_NUMBER() OVER (
                    PARTITION BY {key_list}
                    ORDER BY CASE WHEN data_source = 'release_2025' THEN 0 ELSE 1 END
                ) AS _rn
            FROM {table_name}
            WHERE {null_check}
        )
        SELECT * EXCLUDE (_rn) FROM ranked WHERE _rn = 1
        UNION ALL
        SELECT * FROM {table_name} WHERE {null_check} IS NOT TRUE
    """)

    after = con.execute(f"SELECT COUNT(*) FROM {table_name}").fetchone()[0]
    return before - after


# ---------------------------------------------------------------------------
# Views
# ---------------------------------------------------------------------------

VIEWS = {
    "v_arrest_to_detention": {
        "deps": {"arrests", "detentions"},
        "sql": """
            CREATE OR REPLACE VIEW v_arrest_to_detention AS
            SELECT
                a.*,
                d.stay_book_in_date,
                d.stay_book_out_date,
                d.detention_release_reason,
                d.detention_facility,
                DATEDIFF('day', a.apprehension_date, d.stay_book_in_date)
                    AS days_arrest_to_detention,
                DATEDIFF('day', d.stay_book_in_date, d.stay_book_out_date)
                    AS days_in_detention
            FROM arrests a
            LEFT JOIN detentions d ON a.unique_id = d.unique_id
        """,
    },
    "v_enforcement_pipeline": {
        "deps": {"arrests", "detentions", "removals"},
        "sql": """
            CREATE OR REPLACE VIEW v_enforcement_pipeline AS
            SELECT
                a.unique_id,
                a.apprehension_date,
                a.data_source AS arrest_source,
                d.stay_book_in_date,
                d.stay_book_out_date,
                d.detention_facility,
                r.departure_date,
                r.departure_country AS removal_country,
                DATEDIFF('day', a.apprehension_date, d.stay_book_in_date)
                    AS days_to_detention,
                DATEDIFF('day', d.stay_book_in_date, d.stay_book_out_date)
                    AS days_detained,
                DATEDIFF('day', a.apprehension_date, r.departure_date)
                    AS days_to_removal
            FROM arrests a
            LEFT JOIN detentions d ON a.unique_id = d.unique_id
            LEFT JOIN removals r ON a.unique_id = r.unique_id
        """,
    },
    "v_daily_arrests": {
        "deps": {"arrests"},
        "sql": """
            CREATE OR REPLACE VIEW v_daily_arrests AS
            SELECT
                CAST(apprehension_date AS DATE) AS date,
                data_source,
                COUNT(*) AS arrest_count
            FROM arrests
            WHERE apprehension_date IS NOT NULL
            GROUP BY 1, 2
            ORDER BY 1
        """,
    },
}


def create_views(con: duckdb.DuckDBPyConnection, built: set[str]) -> list[str]:
    created = []
    for name, info in VIEWS.items():
        missing = info["deps"] - built
        if missing:
            print(f"  Skipping {name} (missing: {missing})")
            continue
        try:
            con.execute(info["sql"])
            created.append(name)
            print(f"  Created: {name}")
        except Exception as e:
            print(f"  WARNING: {name}: {e}")
    return created


# ---------------------------------------------------------------------------
# Metadata & validation
# ---------------------------------------------------------------------------


def build_metadata(con: duckdb.DuckDBPyConnection, built: set[str]):
    con.execute("""
        CREATE OR REPLACE TABLE _metadata (
            table_name VARCHAR,
            description VARCHAR,
            row_count BIGINT,
            column_count INTEGER,
            data_sources VARCHAR,
            date_range VARCHAR,
            built_at TIMESTAMP
        )
    """)
    now = datetime.now().isoformat()

    for tbl in sorted(built):
        try:
            rc = con.execute(f"SELECT COUNT(*) FROM {tbl}").fetchone()[0]
            cc = len(con.execute(f"SELECT * FROM {tbl} LIMIT 0").description)

            try:
                srcs = con.execute(
                    f"SELECT DISTINCT data_source FROM {tbl} ORDER BY 1"
                ).fetchall()
                src_str = ", ".join(r[0] for r in srcs)
            except Exception:
                src_str = "unknown"

            dr = ""
            for dc in [
                "apprehension_date", "stay_book_in_date", "departure_date",
                "detainer_prepare_date", "decision_date",
            ]:
                try:
                    r = con.execute(
                        f'SELECT MIN("{dc}")::DATE, MAX("{dc}")::DATE FROM {tbl}'
                    ).fetchone()
                    if r[0]:
                        dr = f"{r[0]} to {r[1]}"
                        break
                except Exception:
                    continue

            desc = TABLE_DESCRIPTIONS.get(tbl, "")
            con.execute(
                "INSERT INTO _metadata VALUES (?,?,?,?,?,?,?)",
                [tbl, desc, rc, cc, src_str, dr, now],
            )
        except Exception as e:
            print(f"  WARNING: metadata for {tbl}: {e}")


def run_validation(con: duckdb.DuckDBPyConnection, db_path: Path):
    print("\n" + "=" * 60)
    print("VALIDATION")
    print("=" * 60)

    try:
        rows = con.execute(
            "SELECT table_name, row_count, column_count, data_sources, date_range "
            "FROM _metadata ORDER BY row_count DESC"
        ).fetchall()
        total = sum(r[1] for r in rows)
        print(f"\n  {len(rows)} tables, {total:,} total rows\n")
        for name, rc, cc, src, dr in rows:
            print(f"    {name:<18s} {rc:>12,} rows  ({cc} cols)  [{src}]  {dr}")
    except Exception as e:
        print(f"  ERROR: {e}")

    print("\n  unique_id coverage:")
    for tbl in ["arrests", "detainers", "detentions", "removals", "rca_decisions"]:
        try:
            tot = con.execute(f"SELECT COUNT(*) FROM {tbl}").fetchone()[0]
            has_id = con.execute(
                f"SELECT COUNT(*) FROM {tbl} WHERE unique_id IS NOT NULL"
            ).fetchone()[0]
            dist = con.execute(
                f"SELECT COUNT(DISTINCT unique_id) FROM {tbl}"
            ).fetchone()[0]
            pct = (has_id / tot * 100) if tot else 0
            print(
                f"    {tbl:<18s} {has_id:>10,}/{tot:>10,} "
                f"({pct:.1f}%)  {dist:>10,} distinct"
            )
        except Exception:
            pass

    print("\n  Rows by data_source:")
    for tbl in ["arrests", "detainers", "detentions", "removals", "rca_decisions"]:
        try:
            rs = con.execute(
                f"SELECT data_source, COUNT(*) FROM {tbl} GROUP BY 1 ORDER BY 1"
            ).fetchall()
            for src, cnt in rs:
                print(f"    {tbl:<18s} {src:<16s} {cnt:>12,}")
        except Exception:
            pass

    print("\n  Cross-table linkage (shared unique_id):")
    for t1, t2 in [
        ("arrests", "detentions"),
        ("arrests", "removals"),
        ("detentions", "removals"),
    ]:
        try:
            cnt = con.execute(f"""
                SELECT COUNT(DISTINCT a.unique_id)
                FROM {t1} a INNER JOIN {t2} b ON a.unique_id = b.unique_id
                WHERE a.unique_id IS NOT NULL
            """).fetchone()[0]
            print(f"    {t1} <-> {t2}: {cnt:,} shared IDs")
        except Exception:
            pass

    size_mb = db_path.stat().st_size / (1024**2)
    print(f"\n  Database size: {size_mb:,.0f} MB ({size_mb/1024:.1f} GB)")


# ---------------------------------------------------------------------------
# Data dictionary
# ---------------------------------------------------------------------------


def build_columns_table(con):
    """Build the _columns data dictionary table."""
    con.execute("DROP TABLE IF EXISTS _columns")

    # Base columns from information_schema
    # ICE _metadata has no source_file column, so we set it to NULL
    con.execute("""
        CREATE TABLE _columns AS
        SELECT
            c.table_name,
            c.column_name,
            c.data_type,
            NULL::VARCHAR AS source_file
        FROM information_schema.columns c
        WHERE c.table_schema = 'main'
          AND c.table_name NOT IN ('_metadata', '_columns')
    """)

    # Add enrichment columns
    con.execute("ALTER TABLE _columns ADD COLUMN example_value VARCHAR")
    con.execute("ALTER TABLE _columns ADD COLUMN join_hint VARCHAR")
    con.execute("ALTER TABLE _columns ADD COLUMN null_pct DOUBLE")

    # Known join hints for ICE
    join_hints = {
        "unique_id": "Anonymized person ID, links across arrests, detainers, detentions, removals",
        "apprehension_date": "Arrest date, joins arrests to detentions by person+date",
        "stay_book_in_date": "Detention start date",
        "departure_date": "Removal/departure date",
        "detention_facility_code": "Facility identifier for detention location",
        "case_id": "EID case identifier",
        "data_source": "Release identifier (release_2023 or release_2025)",
    }

    for col, hint in join_hints.items():
        con.execute(
            "UPDATE _columns SET join_hint = ? WHERE column_name = ?",
            [hint, col],
        )

    # Populate example_value and null_pct for each column
    rows = con.execute(
        "SELECT table_name, column_name FROM _columns"
    ).fetchall()

    for table_name, column_name in rows:
        try:
            result = con.execute(
                f'SELECT CAST("{column_name}" AS VARCHAR) '
                f'FROM "{table_name}" '
                f'WHERE "{column_name}" IS NOT NULL LIMIT 1'
            ).fetchone()
            if result:
                val = result[0]
                if len(val) > 80:
                    val = val[:77] + "..."
                con.execute(
                    "UPDATE _columns SET example_value = ? "
                    "WHERE table_name = ? AND column_name = ?",
                    [val, table_name, column_name],
                )
        except Exception:
            pass

        try:
            result = con.execute(
                f'SELECT ROUND(100.0 * COUNT(*) FILTER (WHERE "{column_name}" IS NULL) '
                f'/ COUNT(*), 1) FROM "{table_name}"'
            ).fetchone()
            if result and result[0] is not None:
                con.execute(
                    "UPDATE _columns SET null_pct = ? "
                    "WHERE table_name = ? AND column_name = ?",
                    [result[0], table_name, column_name],
                )
        except Exception:
            pass


def export_dictionary(con, output_path):
    """Export _columns and _metadata as a readable DICTIONARY.md file."""
    lines = []
    lines.append("# Data Dictionary")
    lines.append("")
    lines.append("Source: [Deportation Data Project](https://deportationdata.org)")
    lines.append("")

    tables = con.execute(
        "SELECT DISTINCT table_name FROM _columns ORDER BY table_name"
    ).fetchall()

    for (table_name,) in tables:
        meta = con.execute(
            "SELECT row_count, description FROM _metadata WHERE table_name = ?",
            [table_name],
        ).fetchone()

        lines.append(f"## {table_name}")
        lines.append("")
        if meta:
            row_count, description = meta
            if description:
                lines.append(f"{description}")
                lines.append("")
            if row_count:
                lines.append(f"Rows: {row_count:,}")
            lines.append("")

        lines.append("| Column | Type | Nulls | Example | Join |")
        lines.append("|--------|------|-------|---------|------|")

        cols = con.execute(
            "SELECT column_name, data_type, null_pct, example_value, join_hint "
            "FROM _columns WHERE table_name = ? ORDER BY rowid",
            [table_name],
        ).fetchall()

        for col_name, dtype, null_pct, example, join_hint in cols:
            null_str = f"{null_pct:.1f}%" if null_pct is not None else ""
            example_str = example if example else ""
            example_str = example_str.replace("|", "\\|")
            join_str = join_hint if join_hint else ""
            lines.append(f"| {col_name} | {dtype} | {null_str} | {example_str} | {join_str} |")

        lines.append("")

    with open(output_path, "w") as f:
        f.write("\n".join(lines))
    print(f"  Exported to {output_path}")


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------


def main():
    parser = argparse.ArgumentParser(
        description="Build DuckDB from ICE enforcement data (Deportation Data Project)"
    )
    parser.add_argument("--data-dir", type=Path, default=Path("data/raw"))
    parser.add_argument("--output", type=Path, default=Path(DEFAULT_OUTPUT))
    parser.add_argument("--tables", nargs="+", help="Build only specific tables")
    args = parser.parse_args()

    t_start = time.time()

    print("=" * 60)
    print("ICE Database Builder (multi-release)")
    print("=" * 60)

    # [1/7] Discover
    print(f"\n[1/7] Discovering files in {args.data_dir}")
    if not args.data_dir.exists():
        print(f"  Error: {args.data_dir} not found")
        sys.exit(1)

    sources = discover_sources(args.data_dir)
    for tbl, chunks in sorted(sources.items()):
        if chunks:
            releases = sorted({get_release_name(f) for f, _ in chunks})
            print(f"    {tbl}: {len(chunks)} sheet(s) from {', '.join(releases)}")

    if args.tables:
        sources = {k: v for k, v in sources.items() if k in args.tables}

    # [2/7] Load
    print(f"\n[2/7] Loading tables")
    db_path = args.output
    db_path.unlink(missing_ok=True)
    con = duckdb.connect(str(db_path))

    built: set[str] = set()
    build_order = ["removals", "rca_decisions", "arrests", "detainers", "detentions"]

    for tbl in build_order:
        chunks = sources.get(tbl, [])
        if not chunks:
            continue
        t0 = time.time()
        rc = load_table(con, tbl, chunks)
        elapsed = time.time() - t0
        if rc > 0:
            built.add(tbl)
            print(f"    -> {tbl}: {rc:,} rows ({elapsed:.1f}s)")

    # [3/7] Dedup
    print(f"\n[3/7] Deduplicating overlapping records")
    for tbl, keys in DEDUP_KEYS.items():
        if tbl not in built:
            continue
        try:
            nsrc = con.execute(
                f"SELECT COUNT(DISTINCT data_source) FROM {tbl}"
            ).fetchone()[0]
            if nsrc <= 1:
                print(f"  {tbl}: single source, skipping")
                continue
        except Exception:
            continue
        removed = dedup_table(con, tbl, keys)
        print(f"  {tbl}: removed {removed:,} duplicates")

    # [4/7] Views
    print(f"\n[4/7] Creating views")
    create_views(con, built)

    # [5/7] Metadata
    print(f"\n[5/7] Building metadata")
    build_metadata(con, built)
    run_validation(con, db_path)

    # [6/7] Build _columns data dictionary
    print(f"\n[6/7] Building _columns data dictionary")
    build_columns_table(con)
    col_count = con.execute("SELECT COUNT(*) FROM _columns").fetchone()[0]
    print(f"  {col_count} columns cataloged in _columns")

    # [7/7] Export DICTIONARY.md
    print(f"\n[7/7] Exporting DICTIONARY.md")
    export_dictionary(con, db_path.parent / "DICTIONARY.md")

    con.close()

    elapsed = time.time() - t_start
    print(f"\nDone in {int(elapsed // 60)}m {int(elapsed % 60)}s")
    print(f"Database: {db_path.resolve()}")


if __name__ == "__main__":
    main()
